#!/usr/bin/env python3
"""Combine the Eventbot CSV data pack into one multi-sheet Excel workbook."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "dashboard_mock_data"
OUTPUT = ROOT / "eventbot_dashboard_all_data.xlsx"

SHEETS = [
    ("Clients", "01_clients.csv"),
    ("Venues", "02_venues.csv"),
    ("Events Calendar", "03_events_calendar.csv"),
    ("Vendor Master", "04_vendor_master.csv"),
    ("Vendor Pricing", "05_vendor_pricing.csv"),
    ("Vendor Schedule", "06_vendor_event_schedule.csv"),
    ("Budget Lines", "07_budget_lines.csv"),
    ("Budget Summary", "08_budget_summary.csv"),
    ("Teams Members", "09_teams_members.csv"),
    ("Event Assignments", "10_event_team_assignments.csv"),
    ("RFQs", "11_rfqs.csv"),
    ("RFQ Items", "12_rfq_items.csv"),
    ("RFQ Quotes", "13_rfq_vendor_quotes.csv"),
    ("RFQ Recipients", "14_rfq_recipients.csv"),
    ("Reputation", "15_reputation_reviews.csv"),
    ("Market Intelligence", "16_market_intelligence.csv"),
    ("Ticket Pricing", "17_ticket_pricing.csv"),
    ("Sponsors", "18_sponsors.csv"),
    ("Wedding Events", "19_wedding_events.csv"),
]

HEADER_FILL = PatternFill("solid", fgColor="131B2E")
HEADER_FONT = Font(color="FFFFFF", bold=True)
LINK_FILL = PatternFill("solid", fgColor="F1F5F9")


def convert(value: str, header: str):
    if value == "":
        return None
    if header.endswith(("_id", "_pct", "_attendees", "_days", "quantity", "inventory", "sold", "remaining", "headcount")):
        try:
            return float(value) if "." in value else int(value)
        except ValueError:
            return value
    if header.endswith("_inr") or header in {"rating", "rating_avg", "reputation_score", "score", "nps", "peak_multiplier"}:
        try:
            return float(value)
        except ValueError:
            return value
    return value


def style_sheet(ws):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 24

    for col_idx, column in enumerate(ws.iter_cols(), 1):
        header = str(column[0].value or "")
        max_width = min(42, max(len(str(c.value or "")) for c in column) + 2)
        ws.column_dimensions[get_column_letter(col_idx)].width = max(11, max_width)
        if header.endswith("_inr"):
            for cell in column[1:]:
                cell.number_format = '₹#,##0.00'
        elif header.endswith("_pct"):
            for cell in column[1:]:
                cell.number_format = '0.00"%"'


def main():
    workbook = Workbook()
    index = workbook.active
    index.title = "Contents"
    index.append(["Sheet", "Source CSV", "Rows", "Description"])

    descriptions = {
        "Clients": "Client contacts and lifetime value",
        "Venues": "Venue capacity, rates, contacts and notes",
        "Events Calendar": "Event dates, IST timings, attendance, revenue and budget",
        "Vendor Master": "Vendor directory, contacts, ratings and reliability",
        "Vendor Pricing": "Rate cards, normal/peak prices, minimum orders and lead times",
        "Vendor Schedule": "Vendor activity linked to event dates, times, venues and RFQs",
        "Budget Lines": "Planned, actual and variance by event/category",
        "Budget Summary": "Event budget utilization, revenue and projected margin",
        "Teams Members": "Teams, people, skills, utilization and ratings",
        "Event Assignments": "Operational team assignments by event",
        "RFQs": "RFQ status, deadlines, ceilings and awards",
        "RFQ Items": "Requested items, quantities and specifications",
        "RFQ Quotes": "Vendor quotes, lead times and ceiling comparisons",
        "RFQ Recipients": "Invitations and response tracking",
        "Reputation": "Reviews, ratings, NPS and sources",
        "Market Intelligence": "Competitors, partners and opportunity signals",
        "Ticket Pricing": "Ticket tiers, inventory, sales and revenue",
        "Sponsors": "Sponsor tiers, amounts and status",
        "Wedding Events": "Wedding-specific calendar, attendance and financial view",
    }

    for sheet_name, filename in SHEETS:
        path = SOURCE / filename
        with path.open(newline="", encoding="utf-8-sig") as handle:
            rows = list(csv.reader(handle))
        ws = workbook.create_sheet(sheet_name)
        headers = rows[0]
        ws.append(headers)
        for row in rows[1:]:
            ws.append([convert(value, headers[i]) for i, value in enumerate(row)])
        style_sheet(ws)
        index.append([sheet_name, filename, max(0, len(rows) - 1), descriptions[sheet_name]])
        index.cell(index.max_row, 1).hyperlink = f"#'{sheet_name}'!A1"
        index.cell(index.max_row, 1).style = "Hyperlink"

    style_sheet(index)
    for row in index.iter_rows(min_row=2):
        for cell in row:
            cell.fill = LINK_FILL
    index.column_dimensions["A"].width = 24
    index.column_dimensions["B"].width = 34
    index.column_dimensions["C"].width = 10
    index.column_dimensions["D"].width = 72
    workbook.save(OUTPUT)
    print(f"Created {OUTPUT} with {len(SHEETS)} data sheets and a Contents sheet")


if __name__ == "__main__":
    main()
