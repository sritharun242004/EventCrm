#!/usr/bin/env python3
"""Create empty CSV templates and one header-only multi-sheet workbook."""

from __future__ import annotations

import csv
import shutil
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "dashboard_mock_data"
PACKAGE = ROOT / "ceo_dashboard_starter"
CSV_DIR = PACKAGE / "empty_csv_templates"
OUTPUT = PACKAGE / "eventbot_empty_dashboard_template.xlsx"


def main():
    CSV_DIR.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    contents = workbook.active
    contents.title = "Contents"
    contents.append(["Sheet", "CSV template", "Purpose"])
    header_fill = PatternFill("solid", fgColor="131B2E")
    header_font = Font(color="FFFFFF", bold=True)

    source_files = sorted(SOURCE.glob("[0-9][0-9]_*.csv"))
    for source in source_files:
        with source.open(newline="", encoding="utf-8-sig") as handle:
            headers = next(csv.reader(handle))
        target = CSV_DIR / source.name
        with target.open("w", newline="", encoding="utf-8") as handle:
            csv.writer(handle).writerow(headers)

        sheet_name = source.stem[3:].replace("_", " ").title()[:31]
        ws = workbook.create_sheet(sheet_name)
        ws.append(headers)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
        for index, header in enumerate(headers, 1):
            ws.column_dimensions[ws.cell(1, index).column_letter].width = max(12, min(30, len(header) + 3))
        contents.append([sheet_name, source.name, "Empty template — add data rows below the header"])
        contents.cell(contents.max_row, 1).hyperlink = f"#'{sheet_name}'!A1"
        contents.cell(contents.max_row, 1).style = "Hyperlink"

    for cell in contents[1]:
        cell.fill = header_fill
        cell.font = header_font
    contents.freeze_panes = "A2"
    contents.column_dimensions["A"].width = 28
    contents.column_dimensions["B"].width = 42
    contents.column_dimensions["C"].width = 52
    workbook.save(OUTPUT)
    shutil.make_archive(str(ROOT / "ceo_dashboard_claude_starter"), "zip", PACKAGE)
    print(f"Created {OUTPUT.name}, {len(source_files)} empty CSV templates, and ceo_dashboard_claude_starter.zip")


if __name__ == "__main__":
    main()
