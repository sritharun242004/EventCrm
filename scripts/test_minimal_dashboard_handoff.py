#!/usr/bin/env python3
"""Simulate the minimal prompt handoff using the distributed mock workbook."""

from __future__ import annotations

import json
import re
from collections import Counter
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "eventbot_dashboard_all_data.xlsx"
STARTER = ROOT / "ceo_dashboard_starter" / "CXO_MINIMAL_DASHBOARD.html"
OUTPUT_DIR = ROOT / "test-results" / "minimal-handoff"
OUTPUT = OUTPUT_DIR / "my_executive_dashboard.html"
TODAY = date(2026, 8, 8)


def rows(sheet):
    values = sheet.iter_rows(values_only=True)
    headers = [str(value) for value in next(values)]
    return [dict(zip(headers, row)) for row in values]


def as_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return datetime.fromisoformat(str(value)).date()


def number(value):
    return float(value or 0)


book = load_workbook(WORKBOOK, read_only=True, data_only=True)
events = rows(book["Events Calendar"])
members = rows(book["Teams Members"])
market = rows(book["Market Intelligence"])

completed = [event for event in events if event["status"] == "completed"]
active_statuses = {"confirmed", "in_production", "live", "wrap_up"}
active = [event for event in events if event["status"] in active_statuses]
upcoming = [event for event in events if as_date(event["starts_on"]) >= TODAY and event["status"] not in {"completed", "cancelled"}]

quarter_data = []
for quarter in range(1, 5):
    quarter_events = [event for event in events if (as_date(event["starts_on"]).month - 1) // 3 + 1 == quarter]
    quarter_data.append({
        "name": f"Q{quarter}",
        "actual": round(sum(number(event["booked_revenue_inr"]) for event in quarter_events)),
        "projected": round(sum(number(event["projected_revenue_inr"]) for event in quarter_events)),
    })

year_actual = round(sum(number(event["booked_revenue_inr"]) for event in events))
year_projected = round(sum(number(event["projected_revenue_inr"]) for event in events))
teams = {row["team_name"] for row in members if row["team_name"]}
people = {row["member_id"] for row in members if row["member_id"]}
utilization = round(sum(number(row["utilization_pct"]) for row in members) / len(members)) if members else None
signal_counts = Counter(str(row["signal"]).lower() for row in market)
pipeline = round(sum(number(event["projected_revenue_inr"]) - number(event["booked_revenue_inr"]) for event in upcoming))
total_budget = round(sum(number(event["total_budget_inr"]) for event in events))
projected_margin = round(100 * (year_projected - total_budget) / year_projected) if year_projected else None
contracted_90d = round(sum(
    number(event["booked_revenue_inr"])
    for event in events
    if TODAY <= as_date(event["starts_on"]) <= TODAY + timedelta(days=90)
))
client_revenue = Counter()
for event in events:
    if event["client"]:
        client_revenue[str(event["client"])] += number(event["booked_revenue_inr"])
top3_concentration = round(100 * sum(value for _, value in client_revenue.most_common(3)) / year_actual) if year_actual else None
overloaded_people = sum(1 for member in members if number(member["utilization_pct"]) > 85)
pipeline_conversion = round(100 * year_actual / year_projected) if year_projected else None

project_rows = sorted(upcoming, key=lambda event: as_date(event["starts_on"]))[:5]
data = {
    "asOf": TODAY.strftime("%d %b %Y"),
    "currency": "₹",
    "financials": {
        "yearActual": year_actual,
        "yearProjected": year_projected,
        "yearBudget": total_budget,
        "outstanding": max(0, year_projected - year_actual),
        "quarters": quarter_data,
    },
    "attention": {
        "cashDue30d": None,
        "overdueCollections": None,
        "projectedMarginPct": projected_margin,
        "overBudgetEvents": None,
        "pipelineConversionPct": pipeline_conversion,
        "contracted90d": contracted_90d,
        "overloadedPeople": overloaded_people,
        "pendingDecisions": None,
        "top3ClientConcentrationPct": top3_concentration,
        "cancellationExposure": None,
    },
    "events": {"completed": len(completed), "inProgress": len(active), "upcoming": len(upcoming), "total": len(events)},
    "teams": {"count": len(teams), "people": len(people), "utilization": utilization},
    "projects": [{
        "name": event["event_name"],
        "client": event["client"] or "Not available",
        "status": str(event["status"]).replace("_", " "),
        "date": as_date(event["starts_on"]).strftime("%d %b"),
        "budget": round(number(event["total_budget_inr"])),
        "collectionPct": round(100 * number(event["booked_revenue_inr"]) / number(event["projected_revenue_inr"])) if number(event["projected_revenue_inr"]) else 0,
    } for event in project_rows],
    "market": {
        "opportunities": signal_counts["opportunity"],
        "competitorEvents": signal_counts["competitor"],
        "partnerSignals": signal_counts["partner"],
        "estimatedPipeline": pipeline,
    },
}

html = STARTER.read_text(encoding="utf-8")
replacement = "const DATA=" + json.dumps(data, ensure_ascii=False, separators=(",", ":")) + ";"
html, changes = re.subn(r"const DATA=\{.*?\};\nconst fmt=", replacement + "\nconst fmt=", html, count=1, flags=re.S)
if changes != 1:
    raise RuntimeError("Could not locate the starter DATA object")

OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
OUTPUT.write_text(html, encoding="utf-8")
print(json.dumps({
    "output": str(OUTPUT),
    "input_sheets_used": ["Events Calendar", "Teams Members", "Market Intelligence"],
    "input_bytes": STARTER.stat().st_size + WORKBOOK.stat().st_size,
    "output_bytes": OUTPUT.stat().st_size,
    "data_object_bytes": len(replacement.encode("utf-8")),
    "events": data["events"],
}, indent=2))
